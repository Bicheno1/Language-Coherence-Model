"""Builds data/vocabulary_sheet.xlsx: the sheet for ASSIGNING axis values to the base vocabulary.

What it contains:
  * vocabulary: the most frequent words of the Brown corpus (lemmas), with what the tables have TODAY for each one
    (type, V/I/Lv/Gv weights, position, cell) and empty YELLOW columns where the new values go. The new position and cell
    are computed with formulas from the new weights.
  * axis_guide: what each axis means, the scale, the formulas, the existing anchors and an example row showing the format.
  * crowded_cells: where nodes pile up (to decide where distinctions are needed).
  * connector_axes: on which axis each relation of the Connectors table falls today (reference for assigning function words).

Priority (column `priority`): 1 function word WITHOUT a node · 2 content word without a node · 3 function with a node (review) ·
4 content with a node (review). Deterministic: same corpus, same tables, same sheet.

Usage:  python tools/build_vocabulary_sheet.py      (needs nltk with 'brown' and 'universal_tagset', and lemminflect)
"""
import os
import sys
from collections import Counter, defaultdict

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "experiments"))
from lcm import load
import exp_geometry as EG                      # reuses the same node anchoring as the experiment
from lemminflect import getLemma
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from nltk.corpus import brown

T = load()
FUNC = {"DET": "determiner", "ADP": "preposition", "PRT": "particle", "CONJ": "conjunction", "PRON": "pronoun"}
CONT = {"NOUN": "noun", "VERB": "verb", "ADJ": "adjective", "ADV": "adverb", "NUM": "numeral"}
AUX = {"be", "have", "do", "will", "would", "can", "could", "may", "might", "must", "shall", "should"}
MIN_FUNC, N_CONT = 30, 1200
# Left out of the sheet: junk tokens from the corpus (tagging/abbreviation leftovers), a historically offensive term, and PROPER NAMES of
# people and places (they are not base axis vocabulary). Proper-name rule: the first WordNet noun sense is an
# instance (washington, kennedy, london…), except common words that WordNet treats that way.
EXCLUDE = {"oth", "af", "aj", "rhode", "negro", "william"}      # william: proper name with no WordNet entry (the rule does not see it)
KEEP_INSTANCES = {"god", "lord", "sun", "earth", "north", "south", "west", "east"}


def is_proper_name(lem):
    from nltk.corpus import wordnet as wn
    ss = wn.synsets(lem, "n")
    return bool(ss and ss[0].instance_hypernyms() and lem not in KEEP_INSTANCES)

ARIAL = "Arial"
F_NORM, F_BOLD = Font(name=ARIAL, size=10), Font(name=ARIAL, size=10, bold=True)
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
FILL_HEAD, FILL_IN, FILL_OUT = PatternFill("solid", fgColor="305496"), PatternFill("solid", fgColor="FFFF00"), PatternFill("solid", fgColor="E2EFDA")


def scan():
    cnt, forms, example = Counter(), defaultdict(Counter), {}
    for s in brown.tagged_sents(tagset="universal"):
        toks, words = [], [w for w, _ in s]
        for w, t in s:
            if not w.isalpha():
                continue
            lw = w.lower()
            if t in CONT and t != "NUM":
                r = getLemma(lw, t)
                lem = r[0].lower() if r else lw
            else:
                lem = lw
            cnt[(lem, t)] += 1
            forms[lem][lw] += 1
            toks.append(lem)
        n = len(words)
        if 6 <= n <= 14:
            text = " ".join(words).replace(" ,", ",").replace(" .", ".").replace(" ;", ";").replace(" ''", '"').replace("`` ", '"')
            for lem in set(toks):
                cur = example.get(lem)
                if cur is None or (8 <= n <= 12 and not 8 <= cur[0] <= 12):
                    example[lem] = (n, text)
    by_lemma = defaultdict(Counter)
    for (lem, t), c in cnt.items():
        by_lemma[lem][t] += c
    total = {lem: sum(pc.values()) for lem, pc in by_lemma.items()}
    dom = {lem: pc.most_common(1)[0][0] for lem, pc in by_lemma.items()}
    return total, dom, forms, example


def select_sample(total, dom):
    func = [lem for lem in total if (dom[lem] in FUNC and total[lem] >= MIN_FUNC and len(lem) > 0) or lem in AUX]
    cont = sorted((lem for lem in total if dom[lem] in CONT and lem not in AUX and len(lem) > 1 and lem.isalpha()
                   and lem not in EXCLUDE and not (dom[lem] == "NOUN" and is_proper_name(lem))),
                  key=lambda w: (-total[w], w))[:N_CONT]
    return sorted(set(func) | set(cont), key=lambda w: (-total[w], w))


def conn_class(lem, pos):
    if lem in AUX:
        return "auxiliary/modal"
    if pos in FUNC:
        return "negation (particle)" if lem in ("not", "n't") else FUNC[pos]
    return CONT.get(pos, pos)


def build():
    print("scanning the corpus…", flush=True)
    total, dom, forms, example = scan()
    lemma_list = select_sample(total, dom)
    print(f"{len(lemma_list)} lemmas selected", flush=True)
    wb = Workbook()

    # ── README (readme sheet) ────────────────────────────────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "README"
    readme = [
        ("ANNOTATION SHEET — LCM BASE VOCABULARY", F_BOLD),
        ("", F_NORM),
        ("What is asked", F_BOLD),
        ("Assign each word four integer weights from 0 to 10 (V, I, Lv, Gv) according to what each axis means. See the 'axis_guide' sheet.", F_NORM),
        ("", F_NORM),
        ("What to edit", F_BOLD),
        ("ONLY the YELLOW cells of the 'vocabulary' sheet: V_new, I_new, Lv_new, Gv_new, criterion and note. The green columns (new x, y, row, col) are formulas: do not touch them.", F_NORM),
        ("Leave empty = no change (the tables' current value is kept). A row counts as assigned only if all four weights have a value.", F_NORM),
        ("", F_NORM),
        ("How it is sorted", F_BOLD),
        ("'rank' is the frequency in the Brown corpus (1 = the most frequent). 'priority': 1 = function word without a node, 2 = content word without a node, 3 = function word with a node (review), 4 = content word with a node (review). Use the filter in row 1.", F_NORM),
        ("", F_NORM),
        ("What it is NOT", F_BOLD),
        ("The 'current' values come from automatic rules and from Unique Host lists: they are a first pass, NOT a truth that must be respected. SentiWordNet and WordNet are not a valid source for these values.", F_NORM),
        ("Left out of the sheet: proper names of people and places (Washington, Kennedy, London…), corpus tagging leftovers and a historically offensive term: they are not base vocabulary. If a word that is missing is needed, add it at the end.", F_NORM),
        ("Words come as LEMMAS (base form); the frequent_forms column shows how they appear in the corpus. corpus_example helps choose the sense.", F_NORM),
        ("", F_NORM),
        ("Data source", F_BOLD),
        ("Frequencies and examples: Brown corpus (NLTK, universal tags), all genres. Current values: data/LCM_tables.json (v0.4i). Generated by tools/build_vocabulary_sheet.py.", F_NORM),
    ]
    for i, (t, f) in enumerate(readme, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    # ── vocabulary (vocabulary sheet) ──────────────────────────────────────────────────────────────────────────────────
    ws = wb.create_sheet("vocabulary")
    cols = ["rank", "word", "class", "pos", "frequency", "frequent_forms", "priority", "in_tables", "table", "type", "category",
            "V", "I", "Lv", "Gv", "x", "y", "row", "col", "corpus_example",
            "V_new", "I_new", "Lv_new", "Gv_new", "x_new", "y_new", "row_new", "col_new", "criterion", "note"]
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    L = {h: get_column_letter(j) for j, h in enumerate(cols, 1)}
    n_missing_f = n_missing_c = n_total = 0
    for r, lem in enumerate(lemma_list, 2):
        pos = dom[lem]
        node = EG.anchor_all(lem, pos)
        cl = conn_class(lem, pos)
        es_func = cl not in CONT.values()
        priority = (1 if es_func else 2) if node is None else (3 if es_func else 4)
        n_total += 1
        n_missing_f += (node is None and es_func)
        n_missing_c += (node is None and not es_func)
        fm = ", ".join(w for w, _ in forms[lem].most_common(3))
        row = [r - 1, lem, cl, pos, total[lem], fm, priority, "yes" if node else "no",
                None, None, None, None, None, None, None, None, None, None, None, example.get(lem, (0, ""))[1]]
        if node:
            tb = next((t for t, n in EG.T.find(node["word"]) if n is node), "")      # EG.T: the instance the node came from
            row[8:19] = [tb, node.get("type"), node.get("category"), node["V"], node["I"], node["Lv"], node["Gv"],
                          node["x"], node["y"], node["row"], node["col"]]
        for j, v in enumerate(row, 1):
            ws.cell(row=r, column=j, value=v).font = F_NORM
        # new columns: yellow weights (to fill in) and green calculation with formulas
        for h in ("V_new", "I_new", "Lv_new", "Gv_new", "criterion", "note"):
            c = ws[f"{L[h]}{r}"]
            c.fill, c.font = FILL_IN, F_NORM
        V, I, Lv, Gv = (f"{L[h]}{r}" for h in ("V_new", "I_new", "Lv_new", "Gv_new"))
        ok = f"COUNT({V}:{Gv})=4"
        ws[f"{L['x_new']}{r}"] = f'=IF({ok},({Lv}-{Gv})/2,"")'
        ws[f"{L['y_new']}{r}"] = f'=IF({ok},({V}-{I})/2,"")'
        ws[f"{L['col_new']}{r}"] = f'=IF({ok},MAX(1,MIN(10,ROUNDUP({L["x_new"]}{r}+5,0))),"")'
        ws[f"{L['row_new']}{r}"] = f'=IF({ok},MAX(1,MIN(10,ROUNDUP(5-{L["y_new"]}{r},0))),"")'
        for h in ("x_new", "y_new", "row_new", "col_new"):
            c = ws[f"{L[h]}{r}"]
            c.fill, c.font = FILL_OUT, F_NORM
    widths = {"rank": 6, "word": 14, "class": 16, "pos": 6, "frequency": 10, "frequent_forms": 22, "priority": 9, "in_tables": 9,
              "table": 10, "type": 14, "category": 14, "corpus_example": 60, "criterion": 30, "note": 30}
    for h, w in widths.items():
        ws.column_dimensions[L[h]].width = w
    for h in ("V", "I", "Lv", "Gv", "x", "y", "row", "col", "V_new", "I_new", "Lv_new", "Gv_new", "x_new", "y_new", "row_new", "col_new"):
        ws.column_dimensions[L[h]].width = 7
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(lemma_list) + 1}"
    ws.row_dimensions[1].height = 30

    # ── axis_guide (axis guide sheet) ─────────────────────────────────────────────────────────────────────────────────────
    g = wb.create_sheet("axis_guide")
    def grid_row(r, vals, bold=False, fill=None):
        for j, v in enumerate(vals, 1):
            c = g.cell(row=r, column=j, value=v)
            c.font = F_BOLD if bold else F_NORM
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if fill:
                c.fill = fill
    r = 1
    grid_row(r, ["Axis (weight)", "Pole", "Side", "Anchor pronoun", "Character", "What goes there"], True, FILL_HEAD); [setattr(c, "font", F_HEAD) for c in g[r]]
    r += 1
    for row_ in [("V", "Context", "up (North)", "We", "Possibility", "future, possibility, condition, purpose; the open, the not yet identified"),
                  ("I", "World", "down (South)", "That", "logical", "past, negation, danger, closure, result; the established, what already exists"),
                  ("Lv", "Imposition", "right (East)", "I", "Sensory", "the sensory, the concrete; the speaker imposes"),
                  ("Gv", "Adaptation", "left (West)", "You", "structural", "the relational, structural, social; the listener; what adapts or receives")]:
        grid_row(r, row_); r += 1
    r += 1
    for t in ["SCALE: integers from 0 to 10 (0 = no weight on that axis).",
              "FORMULAS:  x = (Lv − Gv) / 2    y = (V − I) / 2    column = ceil(x + 5)    row = ceil(5 − y)    (10 × 10 grid; row 1 = Context, row 10 = World; column 1 = Adaptation, column 10 = Imposition).",
              "TENSE: it runs vertically. Past = South (low y); future = North (high y); the PRESENT is the east–west band (y ≈ 0). A past shift is V−2, I+3; a future one V+3, I−2; negation V−3, I+3.",
              "MODE: a question lives in the North (V+2, I−2); an exclamation in the South (V−2, I+2).",
              "PERSPECTIVE: the asker's 'you' is 'I' for the LCM (Imposition) and their 'me' is 'you' (Adaptation).",
              "TENSION = min(V, I): separates the ambivalent (fire V=I=8) from the neutral (stone V=I=1), which the position collapses to the same point.",
              "CURRENT DEFAULT WEIGHTS by type (automatic first pass): biological = V + Gv; non_biological = I + Lv; action = Lv; quality = Gv; plus the per-axis lists from Unique Host and the load of 597 hand-made concepts."]:
        c = g.cell(row=r, column=1, value=t); c.font = F_NORM; c.alignment = Alignment(wrap_text=True, vertical="top")
        g.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        g.row_dimensions[r].height = 44; r += 1
    r += 1
    grid_row(r, ["EXISTING ANCHORS (word)", "V", "I", "Lv", "Gv", "x, y · cell"], True, FILL_HEAD); [setattr(c, "font", F_HEAD) for c in g[r]]
    r += 1
    for w in ["we", "that", "i", "you", "me", "he", "story", "tell", "then", "next", "finally", "because", "but", "however", "if", "and", "who"]:
        hits = [n for _, n in T.find(w)]
        if hits:
            n = hits[0]
            grid_row(r, [w + (f"  [{n.get('class') or n.get('type')}]"), n["V"], n["I"], n["Lv"], n["Gv"], f"({n['x']:+.1f}, {n['y']:+.1f}) · cell {n['row']},{n['col']}"]); r += 1
    r += 1
    grid_row(r, ["PROVISIONAL FUNCTION WORDS (lcm/function_words.py)", "V", "I", "Lv", "Gv", "reason (draft assignment, uncalibrated)"], True, FILL_HEAD); [setattr(c, "font", F_HEAD) for c in g[r]]
    r += 1
    from lcm.function_words import FUNCTION_NODES
    for w, (V, I, Lv, Gv, rationale) in FUNCTION_NODES.items():
        grid_row(r, [w, V, I, Lv, Gv, rationale]); r += 1
    r += 1
    grid_row(r, ["FORMAT EXAMPLE (not an accepted assignment)", "V_new", "I_new", "Lv_new", "Gv_new", "criterion"], True, FILL_HEAD); [setattr(c, "font", F_HEAD) for c in g[r]]
    r += 1
    grid_row(r, ["there", 2, 8, 2, 6, "existential: exists in the world, far from the speaker (World; little Imposition)"], fill=FILL_IN)
    for col, w in zip("ABCDEF", (46, 9, 9, 9, 9, 90)):
        g.column_dimensions[col].width = w

    # ── crowded_cells (crowded cells sheet) ───────────────────────────────────────────────────────────────────────────
    h = wb.create_sheet("crowded_cells")
    for j, t in enumerate(["table", "row", "col", "nodes", "% of the table", "most frequent words in that cell (corpus)"], 1):
        c = h.cell(row=1, column=j, value=t); c.font, c.fill = F_HEAD, FILL_HEAD
    r = 2
    for tb in ("subjects", "concepts"):
        cells = defaultdict(list)
        for n in T[tb].nodes:
            cells[(n["row"], n["col"])].append(n["word"])
        tot = len(T[tb].nodes)
        for (fi, co), ws_ in sorted(cells.items(), key=lambda kv: -len(kv[1]))[:12]:
            top = sorted(ws_, key=lambda w: (-total.get(w.lower(), 0), w))[:10]
            for j, v in enumerate([tb, fi, co, len(ws_), round(len(ws_) / tot, 3), ", ".join(top)], 1):
                h.cell(row=r, column=j, value=v).font = F_NORM
            h.cell(row=r, column=5).number_format = "0.0%"
            r += 1
    for col, w in zip("ABCDEF", (12, 6, 6, 8, 13, 100)):
        h.column_dimensions[col].width = w

    # ── connector_axes (connector axes sheet) ───────────────────────────────────────────────────────────────────────────────
    k = wb.create_sheet("connector_axes")
    for j, t in enumerate(["group", "relation", "current axis (most common)", "no. of connectors", "clases", "examples"], 1):
        c = k.cell(row=1, column=j, value=t); c.font, c.fill = F_HEAD, FILL_HEAD
    grp = defaultdict(list)
    for n in T["connectors"].nodes:
        grp[(n["group"], n["relation"])].append(n)
    r = 2
    for (gr, rel), ns in sorted(grp.items()):
        axes = Counter(n["current_axis"] for n in ns).most_common(1)[0][0]
        for j, v in enumerate([gr, rel, axes, len(ns), ", ".join(sorted({n["class"] for n in ns})), ", ".join(n["word"] for n in ns[:7])], 1):
            k.cell(row=r, column=j, value=v).font = F_NORM
        r += 1
    for col, w in zip("ABCDEF", (14, 36, 22, 14, 18, 70)):
        k.column_dimensions[col].width = w

    out = os.path.join(ROOT, "data", "vocabulary_sheet.xlsx")
    wb.save(out)
    print(f"written: {out}")
    print(f"rows: {n_total} · function words without a node: {n_missing_f} · content words without a node: {n_missing_c}")
    return n_total, n_missing_f, n_missing_c


if __name__ == "__main__":
    build()
